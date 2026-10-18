"""
Match AA Third Party Excel rows to vehicles (procurement_source=third_party)
by normalized chassis and UPDATE specs/prices.

Reads NEXT_PUBLIC_SUPABASE_URL + SUPABASE_SERVICE_ROLE_KEY from .env.local
"""
from __future__ import annotations

import json
import os
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
from supabase import create_client

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SHEETS = [
    Path(r"C:\Users\ATJ\Downloads\AA_11575_2026-10-02_19-55-35_M2fSUp.xlsx"),
    Path(r"C:\Users\ATJ\Downloads\AA_11575_2026-10-02_16-32-57_DssOf6.xlsx"),
]
DEFAULT_TT_RATE = 155
FIXED_EXPENSE_JPY = 45_000


def load_env() -> None:
    for name in (".env.local", ".env"):
        path = ROOT / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip()
            if (v.startswith('"') and v.endswith('"')) or (v.startswith("'") and v.endswith("'")):
                v = v[1:-1]
            os.environ.setdefault(k, v)


def normalize_chassis(raw: object) -> str:
    s = str(raw or "").strip().upper()
    return re.sub(r"[-\s_./\\]", "", s)


def cell_str(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def cell_num(v: object) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[,¥$]", "", str(v)).strip())
    except ValueError:
        return None


def map_transmission(raw: object) -> str | None:
    v = cell_str(raw).lower()
    if not v:
        return None
    if v in {"automatic", "manual", "cvt", "other"}:
        return v
    if v in {"at", "auto", "a/t"}:
        return "automatic"
    if v in {"mt", "manual", "m/t", "stick"}:
        return "manual"
    return "other"


def map_fuel(raw: object) -> str | None:
    v = cell_str(raw).lower()
    if not v:
        return None
    if v in {"petrol", "diesel", "hybrid", "electric", "other"}:
        return v
    if v in {"gasoline", "gas", "ice", "unleaded"}:
        return "petrol"
    if v in {"light oil", "lightoil", "gasoil", "gas oil", "diesel"}:
        return "diesel"
    if v in {"hv", "hev", "phv", "phev"}:
        return "hybrid"
    if v in {"ev", "bev"}:
        return "electric"
    return "other"


def map_status(raw: object) -> str:
    v = cell_str(raw).lower().replace(" ", "_")
    if not v or v in {"available", "stock"}:
        return "in_stock"
    ok = {
        "in_stock",
        "reserved",
        "sold",
        "in_transit",
        "exported",
        "on_hold",
        "archived",
        "hidden",
    }
    return v if v in ok else "in_stock"


def resolve_country(raw: object) -> tuple[str, str]:
    v = cell_str(raw).lower()
    if not v or v in {"japan", "jpn", "jp"}:
        return "JP", "Japan"
    if len(v) == 2:
        return v.upper(), cell_str(raw)
    return "JP", cell_str(raw) or "Japan"


def norm_header(h: object) -> str:
    s = cell_str(h).lower().replace("—", "-").replace("–", "-").replace("−", "-")
    return re.sub(r"\s+", " ", s).strip()


def header_index(headers: list, *aliases: str) -> int:
    wanted = [norm_header(a) for a in aliases]
    # Exact match first (avoid "Model" hitting "Model Code").
    for a in wanted:
        for i, h in enumerate(headers):
            nh = norm_header(h)
            if not nh or "not used" in nh:
                continue
            if nh == a:
                return i
    for a in wanted:
        for i, h in enumerate(headers):
            nh = norm_header(h)
            if not nh or "not used" in nh:
                continue
            if a in nh:
                return i
    return -1


def fob_from_once_jpy(once_price: float) -> dict:
    amount = max(0.0, float(once_price or 0))
    fob_jpy = amount + FIXED_EXPENSE_JPY
    fob_usd = fob_jpy / DEFAULT_TT_RATE
    return {
        "purchase_price": amount,
        "fob_jpy": round(fob_jpy),
        "fob_usd": round(fob_usd),
        "listed_price": round(fob_usd),
    }


def load_sheet_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        matrix = list(ws.iter_rows(values_only=True))
        if len(matrix) < 2:
            continue
        headers = list(matrix[0])
        idx = {
            "chassis": header_index(headers, "chassis no", "chassis", "chassis number"),
            "make": header_index(headers, "make"),
            "model": header_index(headers, "model"),
            "grade": header_index(headers, "version / grade", "version", "grade", "model grade"),
            "year": header_index(headers, "year", "model year"),
            "engine_cc": header_index(headers, "engine cc"),
            "transmission": header_index(headers, "transmission"),
            "color": header_index(headers, "exterior color", "color"),
            "mileage": header_index(headers, "mileage (km)", "mileage"),
            "fuel": header_index(headers, "fuel"),
            "features": header_index(headers, "features"),
            "buying": header_index(headers, "buying price", "once price"),
            "fob": header_index(headers, "fob round", "fob ($)"),
            "status": header_index(headers, "status"),
            "country": header_index(headers, "source country"),
            "port": header_index(headers, "source port"),
            "dealership": header_index(headers, "dealership", "dealership name", "dealer"),
            "lot": header_index(headers, "lot number"),
        }
        if idx["chassis"] < 0:
            print(f"{path.name} / {sheet_name}: no Chassis column — skip")
            continue
        for cells in matrix[1:]:
            cells = list(cells)
            chassis = cell_str(cells[idx["chassis"]] if idx["chassis"] < len(cells) else None)
            if not chassis:
                continue

            def get(key: str):
                i = idx[key]
                if i < 0 or i >= len(cells):
                    return None
                return cells[i]

            rows.append(
                {
                    "source_file": path.name,
                    "sheet": sheet_name,
                    "chassis": chassis,
                    "make": cell_str(get("make")),
                    "model": cell_str(get("model")),
                    "grade": cell_str(get("grade")),
                    "year": cell_num(get("year")),
                    "engine_cc": cell_num(get("engine_cc")),
                    "transmission": get("transmission"),
                    "color": cell_str(get("color")),
                    "mileage": cell_num(get("mileage")),
                    "fuel": get("fuel"),
                    "features": cell_str(get("features")),
                    "buying": cell_num(get("buying")),
                    "fob": cell_num(get("fob")),
                    "status": get("status"),
                    "country": cell_str(get("country")),
                    "port": cell_str(get("port")),
                    "dealership": cell_str(get("dealership")),
                    "lot": cell_str(get("lot")),
                }
            )
    wb.close()
    return rows


def build_payload(row: dict) -> dict:
    code, name = resolve_country(row.get("country") or "Japan")
    once = row.get("buying")
    pricing = fob_from_once_jpy(once) if once is not None else None
    listed_from_sheet = round(row["fob"]) if row.get("fob") is not None else None

    note_bits = []
    if row.get("engine_cc") is not None:
        note_bits.append(f"engine: {int(row['engine_cc'])}cc")
    if row.get("features"):
        note_bits.append(row["features"])
    if row.get("lot"):
        note_bits.append(f"lot: {row['lot']}")
    note_bits.append(f"AA sheet sync {date.today().isoformat()}")

    payload: dict = {
        "make": row["make"] or "Unknown",
        "model": row["model"] or "Pending",
        "model_grade": row["grade"] or None,
        "model_year": int(row["year"]) if row.get("year") is not None else None,
        "color": row["color"] or None,
        "mileage_km": int(row["mileage"]) if row.get("mileage") is not None else None,
        "transmission": map_transmission(row.get("transmission")),
        "fuel_type": map_fuel(row.get("fuel")),
        "source_country": name,
        "source_country_code": code,
        "shipping_port_of_loading": row["port"] or None,
        "status": map_status(row.get("status")),
        "procurement_source": "third_party",
        "notes": "\n".join(note_bits) if note_bits else None,
        "updated_at": date.today().isoformat() + "T00:00:00.000Z",
    }

    if row.get("dealership"):
        payload["dealership_name"] = row["dealership"]
        payload["source_supplier"] = row["dealership"]

    if pricing:
        payload.update(
            {
                "once_price": pricing["purchase_price"],
                "once_price_currency": "JPY",
                "purchase_price": pricing["purchase_price"],
                "purchase_currency": "JPY",
                "expense_jpy": FIXED_EXPENSE_JPY,
                "tt_rate": DEFAULT_TT_RATE,
                "fob_jpy": pricing["fob_jpy"],
                "fob_usd": round(listed_from_sheet)
                if listed_from_sheet is not None
                else pricing["fob_usd"],
                "listed_price": round(listed_from_sheet)
                if listed_from_sheet is not None
                else pricing["listed_price"],
                "listed_currency": "USD",
            }
        )
    elif listed_from_sheet is not None:
        listed = round(listed_from_sheet)
        payload["listed_price"] = listed
        payload["listed_currency"] = "USD"
        payload["fob_usd"] = listed

    return payload


def load_third_party(client) -> list[dict]:
    all_rows: list[dict] = []
    page = 1000
    start = 0
    while True:
        res = (
            client.table("vehicles")
            .select("id, chassis_number, make, model, procurement_source, dealership_name, once_price, listed_price, status")
            .eq("procurement_source", "third_party")
            .range(start, start + page - 1)
            .execute()
        )
        batch = res.data or []
        all_rows.extend(batch)
        if len(batch) < page:
            break
        start += page
    return all_rows


def main() -> int:
    load_env()
    url = (os.environ.get("NEXT_PUBLIC_SUPABASE_URL") or os.environ.get("SUPABASE_URL") or "").strip()
    key = (os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or os.environ.get("SUPABASE_SECRET_KEY") or "").strip()
    if not url or not key:
        print("Missing NEXT_PUBLIC_SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY")
        return 1

    files = [Path(p) for p in sys.argv[1:]] or DEFAULT_SHEETS
    files = [p for p in files if p.exists()]
    if not files:
        print("No Excel files found")
        return 1

    sheet_rows: list[dict] = []
    for f in files:
        rows = load_sheet_rows(f)
        print(f"{f.name}: {len(rows)} data rows")
        sheet_rows.extend(rows)

    by_chassis: dict[str, dict] = {}
    for row in sheet_rows:
        key = normalize_chassis(row["chassis"])
        if key:
            by_chassis[key] = row
    print(f"unique chassis in sheets: {len(by_chassis)}")

    client = create_client(url, key)
    vehicles = load_third_party(client)
    print(f"third_party vehicles in DB: {len(vehicles)}")

    by_norm = {normalize_chassis(v["chassis_number"]): v for v in vehicles if normalize_chassis(v.get("chassis_number"))}

    matched = updated = failed = 0
    unmatched_sheet: list[str] = []
    unmatched_db: list[str] = []
    errors: list[str] = []

    for norm, row in by_chassis.items():
        vehicle = by_norm.get(norm)
        if not vehicle:
            unmatched_sheet.append(row["chassis"])
            continue
        matched += 1
        payload = build_payload(row)
        try:
            client.table("vehicles").update(payload).eq("id", vehicle["id"]).execute()
            updated += 1
            by_norm.pop(norm, None)
        except Exception as exc:  # noqa: BLE001
            failed += 1
            errors.append(f"{row['chassis']}: {exc}")

    for v in by_norm.values():
        unmatched_db.append(v["chassis_number"])

    summary = {
        "sheet_rows": len(sheet_rows),
        "unique_sheet_chassis": len(by_chassis),
        "db_third_party": len(vehicles),
        "matched": matched,
        "updated": updated,
        "failed": failed,
        "unmatched_sheet": len(unmatched_sheet),
        "unmatched_db_third_party": len(unmatched_db),
    }
    print("\n=== sync summary ===")
    print(json.dumps(summary, indent=2))
    if unmatched_sheet:
        print("sample unmatched sheet chassis:", unmatched_sheet[:25])
    if unmatched_db:
        print("sample unmatched DB third_party chassis:", unmatched_db[:25])
    if errors:
        print("errors:", errors[:20])
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
